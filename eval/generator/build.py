"""Build driver: write all xlsx files and their ground-truth label JSON.

    python -m eval.generator.build

Steps:
  1. spec-driven workbooks (static values)  -> render + labels
  2. extreme stress-test workbooks (live formulas / cross-sheet refs / named ranges)
  3. recalculate the formula-bearing files with LibreOffice headless so they carry
     cached results (a fair test: a real saved file would have them too)
  4. write label JSON
"""
from __future__ import annotations

import json
import random
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook

from eval.generator import hard_workbooks, scale_workbook
from eval.generator.sampling import make_samples, prune_cells, resolve_measures
from eval.generator.specs import WorkbookSpec, all_specs
from eval.generator.tables import render_table
from eval.schemas import TableLabel, WorkbookLabel

HERE = Path(__file__).resolve().parents[1]
DATA = HERE / "data"
WB_DIR = DATA / "workbooks"
LB_DIR = DATA / "labels"


def build_spec(spec: WorkbookSpec) -> WorkbookLabel:
    wb = Workbook()
    wb.remove(wb.active)
    sheet_order: list[str] = []
    for t in spec.tables:
        if t.sheet not in sheet_order:
            sheet_order.append(t.sheet)
    sheets = {name: wb.create_sheet(title=name) for name in sheet_order}

    tables: list[TableLabel] = [render_table(sheets[t.sheet], t) for t in spec.tables]
    measures = resolve_measures(tables, spec.measures)
    samples, referenced = make_samples(
        spec.filename, tables, measures, spec.formulas, spec.business_logic)
    prune_cells(tables, referenced, random.Random(hash(spec.filename) & 0xFFFF))

    label = WorkbookLabel(
        workbook=spec.filename, rel_path=f"workbooks/{spec.filename}",
        difficulty=spec.difficulty, domain=spec.domain, traps=spec.traps,
        sheets=sheet_order, business_logic=spec.business_logic,
        tables=tables, measures=measures, samples=samples)
    WB_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(WB_DIR / spec.filename)
    return label


def recalc_with_libreoffice(filenames: list[str]) -> None:
    """Open each file in LibreOffice headless and re-save so formulas get cached."""
    if not filenames:
        return
    paths = [str(WB_DIR / f) for f in filenames]
    with tempfile.TemporaryDirectory() as tmp:
        cmd = [
            "soffice", "--headless", "--calc", "--convert-to", "xlsx",
            "--outdir", tmp,
            "-env:UserInstallation=file:///tmp/lo_eval_profile",
            *paths,
        ]
        subprocess.run(cmd, check=True, capture_output=True, timeout=180)
        for f in filenames:
            out = Path(tmp) / f
            if out.exists():
                shutil.move(str(out), str(WB_DIR / f))


def write_label(label: WorkbookLabel) -> None:
    LB_DIR.mkdir(parents=True, exist_ok=True)
    (LB_DIR / f"{Path(label.workbook).stem}.json").write_text(
        label.model_dump_json(indent=2))


def _print(label: WorkbookLabel) -> None:
    by_type: dict[str, int] = {}
    for s in label.samples:
        by_type[s.type] = by_type.get(s.type, 0) + 1
    n_formula_cells = sum(1 for t in label.tables for c in t.cells if c.is_formula)
    extra = f" formula_cells={n_formula_cells}" if n_formula_cells else ""
    print(f"  {label.workbook:36s} {label.difficulty:7s} "
          f"tables={len(label.tables)} measures={len(label.measures)} "
          f"samples={len(label.samples)} {by_type}{extra}")


def main() -> None:
    labels: list[WorkbookLabel] = []

    for spec in all_specs():
        labels.append(build_spec(spec))

    recalc_files: list[str] = []
    for label, needs_recalc in hard_workbooks.build_all(WB_DIR):
        labels.append(label)
        if needs_recalc:
            recalc_files.append(label.workbook)

    scale_label, _ = scale_workbook.build(WB_DIR)
    labels.append(scale_label)

    recalc_with_libreoffice(recalc_files)

    for label in labels:
        write_label(label)
        _print(label)

    total = sum(len(l.samples) for l in labels)
    print(f"\nBuilt {len(labels)} workbooks, {total} validation samples "
          f"({len(recalc_files)} recalculated with LibreOffice).")


if __name__ == "__main__":
    main()

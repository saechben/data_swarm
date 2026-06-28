"""Task 2: test that split_workbook and run_swarm accept WorkbookSource objects."""
import openpyxl
from mcg_swarm.source import OpenpyxlFileSource
from mcg_swarm.runner import run_swarm
from mcg_swarm.splitter import split_workbook


def _wb(tmp_path):
    p = tmp_path / "p.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Revenue"])
    ws.append(["NA", 10])
    ws.append(["EU", 20])
    wb.save(p)
    return str(p)


def test_split_workbook_accepts_source(tmp_path):
    src = OpenpyxlFileSource(_wb(tmp_path))
    handles = split_workbook(src)
    assert [h.sheet for h in handles] == ["Data"]
    assert [c.name for c in handles[0].columns] == ["Region", "Revenue"]


def test_run_swarm_accepts_path_and_source(tmp_path):
    p = _wb(tmp_path)
    a = run_swarm({"main": p})           # back-compat dict
    b = run_swarm(OpenpyxlFileSource(p)) # explicit source
    assert a.tables[0].columns[0].name == b.tables[0].columns[0].name == "Region"

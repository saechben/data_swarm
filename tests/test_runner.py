import openpyxl
from mcg_swarm.runner import run_swarm, build_indices


def _wb_two_tabs(tmp_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = "Sales"
    for r in [["Region", "Revenue"], ["EMEA", 100], ["APAC", 200]]: ws1.append(r)
    ws2 = wb.create_sheet("Empty")  # ambiguous tab
    p = tmp_path / "multi.xlsx"; wb.save(p); return str(p)


def test_run_swarm_one_table_per_tab_bad_tab_isolated(tmp_path):
    p = _wb_two_tabs(tmp_path)
    ext = run_swarm({"main": p})
    assert ext.workbook.endswith("multi.xlsx")
    assert len(ext.tables) == 2
    sales = [t for t in ext.tables if t.sheet == "Sales"][0]
    empty = [t for t in ext.tables if t.sheet == "Empty"][0]
    assert sales.errors == [] and empty.errors  # bad tab isolated, file still processed


def test_build_indices_round_trips(tmp_path):
    p = _wb_two_tabs(tmp_path)
    ext = run_swarm({"main": p})
    idxs = build_indices(p, ext)
    sales = [t for t in ext.tables if t.sheet == "Sales"][0]
    assert idxs[sales.table_id].query("APAC", "Revenue").value == 200


def test_capex_gate_no_crash():
    """E2E guard: capex_plan.xlsx has banner/offset empty cells that previously caused
    EmptyCell AttributeError in run_table_tests → orchestrate_table caught it → errors set.
    Post-fix: the table must have errors==[] and its index must include expected columns.
    """
    import os
    capex_path = os.path.join(
        os.path.dirname(__file__), "..", "eval", "data", "workbooks", "capex_plan.xlsx"
    )
    capex_path = os.path.normpath(capex_path)
    if not os.path.exists(capex_path):
        import pytest
        pytest.skip("capex_plan.xlsx not found — skipping E2E guard")

    ext = run_swarm({"main": capex_path}, llm=None)
    capex_tables = [t for t in ext.tables if "capex" in t.table_id.lower() or t.sheet.lower() == "capex"]
    assert capex_tables, f"No Capex table found; tables: {[t.table_id for t in ext.tables]}"
    capex = capex_tables[0]
    assert capex.errors == [], f"Capex gate crashed with errors: {capex.errors}"
    col_names = [c.name for c in capex.columns]
    assert any("2026" in n or "Total" in n for n in col_names), (
        f"Expected Y2026/Total columns in capex, got: {col_names}"
    )
